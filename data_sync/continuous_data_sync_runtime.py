from __future__ import annotations

import csv
import hashlib
import json
import sqlite3
import uuid
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any


SUPPORTED_SUFFIXES = {
    ".csv",
    ".json",
    ".jsonl",
    ".xlsx",
    ".docx",
    ".pdf",
    ".txt",
    ".jpg",
    ".jpeg",
    ".png",
    ".webp",
    ".tiff",
}

RUNTIME_CONTRACT_VERSION = "1.0"
SYNC_SCHEMA_SQL = """
CREATE TABLE IF NOT EXISTS sync_batches (
  batch_id TEXT PRIMARY KEY, tenant_id TEXT NOT NULL, project_id TEXT NOT NULL,
  source_id TEXT NOT NULL, status TEXT NOT NULL, manifest_hash TEXT NOT NULL,
  started_at TEXT NOT NULL, completed_at TEXT, published_version INTEGER, errors_json TEXT NOT NULL DEFAULT '[]'
);
CREATE TABLE IF NOT EXISTS candidate_records (
  batch_id TEXT NOT NULL, tenant_id TEXT NOT NULL, project_id TEXT NOT NULL, source_id TEXT NOT NULL,
  record_key TEXT NOT NULL, payload_json TEXT NOT NULL, source_path TEXT NOT NULL,
  content_checksum TEXT NOT NULL, deleted INTEGER NOT NULL DEFAULT 0,
  PRIMARY KEY (batch_id, record_key)
);
CREATE TABLE IF NOT EXISTS publish_versions (
  tenant_id TEXT NOT NULL, project_id TEXT NOT NULL, source_id TEXT NOT NULL,
  version INTEGER NOT NULL, batch_id TEXT NOT NULL, manifest_hash TEXT NOT NULL, published_at TEXT NOT NULL,
  PRIMARY KEY (tenant_id, project_id, source_id, version)
);
CREATE TABLE IF NOT EXISTS published_records (
  tenant_id TEXT NOT NULL, project_id TEXT NOT NULL, source_id TEXT NOT NULL, version INTEGER NOT NULL,
  record_key TEXT NOT NULL, payload_json TEXT NOT NULL, source_path TEXT NOT NULL,
  content_checksum TEXT NOT NULL, deleted INTEGER NOT NULL DEFAULT 0,
  PRIMARY KEY (tenant_id, project_id, source_id, version, record_key)
);
CREATE TABLE IF NOT EXISTS current_versions (
  tenant_id TEXT NOT NULL, project_id TEXT NOT NULL, source_id TEXT NOT NULL,
  version INTEGER NOT NULL, switched_at TEXT NOT NULL,
  PRIMARY KEY (tenant_id, project_id, source_id)
);
CREATE VIEW IF NOT EXISTS current_published_records AS
SELECT r.* FROM published_records r
JOIN current_versions c ON c.tenant_id=r.tenant_id AND c.project_id=r.project_id
  AND c.source_id=r.source_id AND c.version=r.version
WHERE r.deleted=0;
CREATE TABLE IF NOT EXISTS sync_jobs (
  job_id TEXT PRIMARY KEY, tenant_id TEXT NOT NULL, project_id TEXT NOT NULL, source_id TEXT NOT NULL,
  dedupe_key TEXT NOT NULL, state TEXT NOT NULL, attempts INTEGER NOT NULL DEFAULT 0,
  result_json TEXT NOT NULL DEFAULT '{}', error_code TEXT NOT NULL DEFAULT '',
  created_at TEXT NOT NULL, updated_at TEXT NOT NULL,
  UNIQUE(tenant_id, project_id, source_id, dedupe_key)
);
CREATE INDEX IF NOT EXISTS idx_sync_jobs_claim ON sync_jobs(state, created_at);
CREATE TABLE IF NOT EXISTS rollback_events (
  id INTEGER PRIMARY KEY AUTOINCREMENT, tenant_id TEXT NOT NULL, project_id TEXT NOT NULL,
  source_id TEXT NOT NULL, from_version INTEGER, to_version INTEGER NOT NULL,
  idempotency_key TEXT NOT NULL, actor_hash TEXT NOT NULL, created_at TEXT NOT NULL,
  UNIQUE(tenant_id, project_id, source_id, idempotency_key)
);
"""


@dataclass(frozen=True)
class SyncScope:
    tenant_id: str
    project_id: str
    source_id: str


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def run_directory_sync(
    *,
    contract: dict[str, Any],
    source_dir: Path,
    database_path: Path,
) -> dict[str, Any]:
    """Synchronize one local directory through candidate validation and atomic publish.

    This runtime is deliberately provider-neutral. Cloud credentials and schedules are
    attached by deployment adapters; this function proves the data lifecycle itself.
    """
    scope = _scope(contract)
    source = Path(source_dir).expanduser().resolve()
    database = Path(database_path).expanduser().resolve()
    if not source.is_dir():
        raise FileNotFoundError(f"source directory not found: {source}")
    database.parent.mkdir(parents=True, exist_ok=True)

    files = _safe_source_files(source, contract)
    batch_id = uuid.uuid4().hex
    manifest = [_file_manifest(path, source) for path in files]
    manifest_hash = _hash_json({
        "runtime_contract_version": RUNTIME_CONTRACT_VERSION,
        "contract": contract,
        "files": manifest,
    })
    parsed_records: list[dict[str, Any]] = []
    errors: list[dict[str, str]] = []
    for path in files:
        try:
            parsed_records.extend(_parse_file(path, source, contract))
        except Exception as exc:
            errors.append({"path": path.relative_to(source).as_posix(), "error": str(exc)})
    duplicate_keys = _duplicate_record_keys(parsed_records)
    errors.extend({"path": "record_identity", "error": f"duplicate record_key: {key}"} for key in duplicate_keys)

    with sqlite3.connect(database) as connection:
        connection.row_factory = sqlite3.Row
        _ensure_schema(connection)
        _record_batch_start(connection, scope, batch_id, manifest_hash)
        if errors:
            _finish_batch(connection, batch_id, "completed_with_mismatch", None, errors)
            connection.commit()
            return _result(scope, batch_id, "completed_with_mismatch", manifest_hash, None, parsed_records, errors, database)
        _replace_candidate(connection, scope, batch_id, parsed_records)

        current = _current_version(connection, scope)
        if current and current["manifest_hash"] == manifest_hash:
            _finish_batch(connection, batch_id, "verified_no_change", int(current["version"]), [])
            connection.commit()
            return _result(scope, batch_id, "verified_no_change", manifest_hash, int(current["version"]), parsed_records, [], database)

        version = (int(current["version"]) if current else 0) + 1
        _publish_atomically(connection, scope, batch_id, version, manifest_hash, parsed_records)
        _finish_batch(connection, batch_id, "verified", version, [])
        connection.commit()
    return _result(scope, batch_id, "verified", manifest_hash, version, parsed_records, [], database)


def read_published(*, database_path: Path, tenant_id: str, project_id: str, source_id: str) -> list[dict[str, Any]]:
    scope = SyncScope(tenant_id, project_id, source_id)
    with sqlite3.connect(Path(database_path)) as connection:
        connection.row_factory = sqlite3.Row
        current = _current_version(connection, scope)
        if not current:
            return []
        rows = connection.execute(
            """
            SELECT record_key, payload_json, source_path, content_checksum, deleted
            FROM published_records
            WHERE tenant_id=? AND project_id=? AND source_id=? AND version=?
            ORDER BY record_key
            """,
            (*_scope_values(scope), int(current["version"])),
        ).fetchall()
    return [
        {
            "record_key": row["record_key"],
            "payload": json.loads(row["payload_json"]),
            "source_path": row["source_path"],
            "content_checksum": row["content_checksum"],
            "deleted": bool(row["deleted"]),
        }
        for row in rows
        if not row["deleted"]
    ]


def rollback_to_version(
    *,
    database_path: Path,
    tenant_id: str,
    project_id: str,
    source_id: str,
    version: int,
    idempotency_key: str = "",
    actor_hash: str = "system",
) -> dict[str, Any]:
    scope = SyncScope(tenant_id, project_id, source_id)
    with sqlite3.connect(Path(database_path)) as connection:
        _ensure_schema(connection)
        exists = connection.execute(
            "SELECT 1 FROM publish_versions WHERE tenant_id=? AND project_id=? AND source_id=? AND version=?",
            (*_scope_values(scope), version),
        ).fetchone()
        if not exists:
            raise ValueError(f"publish version does not exist: {version}")
        current = connection.execute(
            "SELECT version FROM current_versions WHERE tenant_id=? AND project_id=? AND source_id=?",
            _scope_values(scope),
        ).fetchone()
        if idempotency_key:
            prior = connection.execute(
                "SELECT to_version FROM rollback_events WHERE tenant_id=? AND project_id=? AND source_id=? AND idempotency_key=?",
                (*_scope_values(scope), idempotency_key),
            ).fetchone()
            if prior:
                return {"status": "already_rolled_back", "version": int(prior[0]), **scope.__dict__}
        connection.execute(
            """
            UPDATE current_versions SET version=?, switched_at=?
            WHERE tenant_id=? AND project_id=? AND source_id=?
            """,
            (version, utc_now(), *_scope_values(scope)),
        )
        if idempotency_key:
            connection.execute(
                "INSERT INTO rollback_events(tenant_id,project_id,source_id,from_version,to_version,idempotency_key,actor_hash,created_at) VALUES(?,?,?,?,?,?,?,?)",
                (*_scope_values(scope), int(current[0]) if current else None, version, idempotency_key, actor_hash, utc_now()),
            )
        connection.commit()
    return {"status": "rolled_back", "version": version, **scope.__dict__}


class SyncJobRuntime:
    """Durable, project-scoped queue for bounded Stage 6 worker execution."""

    def __init__(self, database_path: Path | str, *, max_attempts: int = 3) -> None:
        self.database_path = Path(database_path)
        self.database_path.parent.mkdir(parents=True, exist_ok=True)
        self.max_attempts = max(1, int(max_attempts))
        with sqlite3.connect(self.database_path) as connection:
            _ensure_schema(connection)

    def enqueue(self, *, contract: dict[str, Any], dedupe_key: str) -> dict[str, Any]:
        scope = _scope(contract)
        key = str(dedupe_key).strip()
        if not key:
            return {"status": "blocked", "reason": "dedupe_key_required"}
        job_id = uuid.uuid4().hex
        timestamp = utc_now()
        try:
            with sqlite3.connect(self.database_path) as connection:
                _ensure_schema(connection)
                connection.execute(
                    "INSERT INTO sync_jobs(job_id,tenant_id,project_id,source_id,dedupe_key,state,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?)",
                    (job_id, *_scope_values(scope), key, "pending", timestamp, timestamp),
                )
            return {"status": "enqueued", "job_id": job_id, **scope.__dict__}
        except sqlite3.IntegrityError:
            return {"status": "duplicate", "reason": "dedupe_key_already_exists", **scope.__dict__}

    def run_next(
        self,
        *,
        contract: dict[str, Any],
        source_dir: Path,
        published_database_path: Path,
    ) -> dict[str, Any]:
        scope = _scope(contract)
        timestamp = utc_now()
        with sqlite3.connect(self.database_path) as connection:
            connection.row_factory = sqlite3.Row
            _ensure_schema(connection)
            connection.execute("BEGIN IMMEDIATE")
            row = connection.execute(
                "SELECT * FROM sync_jobs WHERE tenant_id=? AND project_id=? AND source_id=? AND state IN ('pending','retry') ORDER BY created_at,job_id LIMIT 1",
                _scope_values(scope),
            ).fetchone()
            if row is None:
                connection.commit()
                return {"status": "idle", **scope.__dict__}
            attempts = int(row["attempts"]) + 1
            connection.execute(
                "UPDATE sync_jobs SET state='running',attempts=?,updated_at=? WHERE job_id=?",
                (attempts, timestamp, row["job_id"]),
            )
            connection.commit()
        try:
            result = run_directory_sync(contract=contract, source_dir=source_dir, database_path=published_database_path)
            state = "completed" if result["status"] in {"verified", "verified_no_change"} else "blocked_mismatch"
            error_code = "" if state == "completed" else "candidate_validation_failed"
        except Exception as exc:
            result = {"status": "failed", "error_type": type(exc).__name__}
            state = "dead_letter" if attempts >= self.max_attempts else "retry"
            error_code = type(exc).__name__[:80]
        with sqlite3.connect(self.database_path) as connection:
            connection.execute(
                "UPDATE sync_jobs SET state=?,result_json=?,error_code=?,updated_at=? WHERE job_id=? AND state='running'",
                (state, json.dumps(result, ensure_ascii=False, sort_keys=True), error_code, utc_now(), row["job_id"]),
            )
        return {"status": state, "job_id": row["job_id"], "attempts": attempts, "result": result, **scope.__dict__}

    def status(self, *, tenant_id: str, project_id: str, source_id: str) -> dict[str, Any]:
        with sqlite3.connect(self.database_path) as connection:
            _ensure_schema(connection)
            rows = connection.execute(
                "SELECT state,COUNT(*) FROM sync_jobs WHERE tenant_id=? AND project_id=? AND source_id=? GROUP BY state",
                (tenant_id, project_id, source_id),
            ).fetchall()
        return {"status": "ok", "counts": {str(state): int(count) for state, count in rows}}


def _scope(contract: dict[str, Any]) -> SyncScope:
    values = [str(contract.get(key, "")).strip() for key in ["tenant_id", "project_id", "source_id"]]
    if not all(values):
        raise ValueError("tenant_id, project_id and source_id are required")
    return SyncScope(*values)


def _safe_source_files(source: Path, contract: dict[str, Any]) -> list[Path]:
    max_files = int(contract.get("limits", {}).get("max_files_per_batch", 1000))
    max_bytes = int(contract.get("limits", {}).get("max_file_bytes", 25 * 1024 * 1024))
    files: list[Path] = []
    for path in sorted(source.rglob("*")):
        if not path.is_file() or path.is_symlink():
            continue
        resolved = path.resolve()
        if source not in resolved.parents:
            raise ValueError(f"source path escapes root: {path}")
        if path.suffix.lower() not in SUPPORTED_SUFFIXES:
            continue
        if path.stat().st_size > max_bytes:
            raise ValueError(f"file exceeds max_file_bytes: {path.name}")
        files.append(path)
        if len(files) > max_files:
            raise ValueError("source exceeds max_files_per_batch")
    return files


def _file_manifest(path: Path, source: Path) -> dict[str, Any]:
    data = path.read_bytes()
    return {
        "path": path.relative_to(source).as_posix(),
        "suffix": path.suffix.lower(),
        "size": len(data),
        "checksum": hashlib.sha256(data).hexdigest(),
    }


def _parse_file(path: Path, source: Path, contract: dict[str, Any]) -> list[dict[str, Any]]:
    suffix = path.suffix.lower()
    relative = path.relative_to(source).as_posix()
    checksum = hashlib.sha256(path.read_bytes()).hexdigest()
    if suffix == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as handle:
            rows = list(csv.DictReader(handle))
        return _rows_to_records(rows, relative, checksum, contract)
    if suffix == ".json":
        payload = json.loads(path.read_text(encoding="utf-8"))
        rows = payload if isinstance(payload, list) else [payload]
        if not all(isinstance(row, dict) for row in rows):
            raise ValueError("JSON must contain an object or list of objects")
        return _rows_to_records(rows, relative, checksum, contract)
    if suffix == ".jsonl":
        rows = [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]
        if not all(isinstance(row, dict) for row in rows):
            raise ValueError("JSONL lines must be objects")
        return _rows_to_records(rows, relative, checksum, contract)
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        rows: list[dict[str, Any]] = []
        for sheet in workbook.worksheets:
            iterator = sheet.iter_rows(values_only=True)
            headers = [str(value or "").strip() for value in next(iterator, [])]
            if not any(headers):
                continue
            for values in iterator:
                row = dict(zip(headers, values))
                row["_source_sheet"] = sheet.title
                rows.append(row)
        return _rows_to_records(rows, relative, checksum, contract)
    if suffix == ".docx":
        from docx import Document

        document = Document(path)
        text = "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())
        tables = [[[cell.text for cell in row.cells] for row in table.rows] for table in document.tables]
        return [_asset_record(relative, checksum, {"kind": "docx", "text": text, "tables": tables})]
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(path)
        pages = [page.extract_text() or "" for page in reader.pages]
        if not any(value.strip() for value in pages) and contract.get("document_policy", {}).get("scanned_pdf") == "quarantine_without_ocr":
            raise ValueError("scanned PDF requires an authorized OCR provider")
        return [_asset_record(relative, checksum, {"kind": "pdf", "pages": pages, "page_count": len(pages)})]
    if suffix in {".jpg", ".jpeg", ".png", ".webp", ".tiff"}:
        from PIL import Image

        with Image.open(path) as image:
            metadata = {"kind": "image_asset", "width": image.width, "height": image.height, "format": image.format}
        if contract.get("document_policy", {}).get("image_text") == "ocr_required":
            raise ValueError("image text extraction requires an authorized OCR provider")
        return [_asset_record(relative, checksum, metadata)]
    return [_asset_record(relative, checksum, {"kind": "text", "text": path.read_text(encoding="utf-8")})]


def _rows_to_records(rows: list[dict[str, Any]], relative: str, checksum: str, contract: dict[str, Any]) -> list[dict[str, Any]]:
    identity_fields = [str(value) for value in contract.get("record_identity", []) if str(value).strip()]
    records = []
    for index, row in enumerate(rows, start=1):
        cleaned = {str(key): value for key, value in row.items()}
        identity = [str(cleaned.get(field, "")).strip() for field in identity_fields]
        key = "|".join(identity) if identity and all(identity) else f"{relative}#{index}"
        records.append({
            "record_key": key,
            "payload": cleaned,
            "source_path": relative,
            "content_checksum": _hash_json(cleaned),
            "deleted": bool(cleaned.get("_deleted", False)),
        })
    return records


def _asset_record(relative: str, checksum: str, payload: dict[str, Any]) -> dict[str, Any]:
    return {"record_key": relative, "payload": payload, "source_path": relative, "content_checksum": checksum, "deleted": False}


def _ensure_schema(connection: sqlite3.Connection) -> None:
    connection.executescript(SYNC_SCHEMA_SQL)


def _record_batch_start(connection: sqlite3.Connection, scope: SyncScope, batch_id: str, manifest_hash: str) -> None:
    connection.execute(
        "INSERT INTO sync_batches(batch_id,tenant_id,project_id,source_id,status,manifest_hash,started_at) VALUES(?,?,?,?,?,?,?)",
        (batch_id, *_scope_values(scope), "candidate", manifest_hash, utc_now()),
    )


def _replace_candidate(connection: sqlite3.Connection, scope: SyncScope, batch_id: str, records: list[dict[str, Any]]) -> None:
    for record in records:
        connection.execute(
            """
            INSERT INTO candidate_records(batch_id,tenant_id,project_id,source_id,record_key,payload_json,source_path,content_checksum,deleted)
            VALUES(?,?,?,?,?,?,?,?,?)
            """,
            (batch_id, *_scope_values(scope), record["record_key"], json.dumps(record["payload"], ensure_ascii=False, sort_keys=True, default=str), record["source_path"], record["content_checksum"], int(record["deleted"])),
        )


def _current_version(connection: sqlite3.Connection, scope: SyncScope) -> sqlite3.Row | None:
    return connection.execute(
        """
        SELECT c.version, p.manifest_hash FROM current_versions c
        JOIN publish_versions p ON p.tenant_id=c.tenant_id AND p.project_id=c.project_id
          AND p.source_id=c.source_id AND p.version=c.version
        WHERE c.tenant_id=? AND c.project_id=? AND c.source_id=?
        """,
        _scope_values(scope),
    ).fetchone()


def _publish_atomically(connection: sqlite3.Connection, scope: SyncScope, batch_id: str, version: int, manifest_hash: str, records: list[dict[str, Any]]) -> None:
    previous = _current_version(connection, scope)
    snapshot: dict[str, dict[str, Any]] = {}
    if previous:
        rows = connection.execute(
            "SELECT record_key,payload_json,source_path,content_checksum,deleted FROM published_records WHERE tenant_id=? AND project_id=? AND source_id=? AND version=?",
            (*_scope_values(scope), int(previous["version"])),
        ).fetchall()
        snapshot = {row[0]: {"record_key": row[0], "payload": json.loads(row[1]), "source_path": row[2], "content_checksum": row[3], "deleted": bool(row[4])} for row in rows}
    for record in records:
        snapshot[record["record_key"]] = record
    connection.execute(
        "INSERT INTO publish_versions(tenant_id,project_id,source_id,version,batch_id,manifest_hash,published_at) VALUES(?,?,?,?,?,?,?)",
        (*_scope_values(scope), version, batch_id, manifest_hash, utc_now()),
    )
    for record in snapshot.values():
        connection.execute(
            """
            INSERT INTO published_records(tenant_id,project_id,source_id,version,record_key,payload_json,source_path,content_checksum,deleted)
            VALUES(?,?,?,?,?,?,?,?,?)
            """,
            (*_scope_values(scope), version, record["record_key"], json.dumps(record["payload"], ensure_ascii=False, sort_keys=True, default=str), record["source_path"], record["content_checksum"], int(record["deleted"])),
        )
    connection.execute(
        """
        INSERT INTO current_versions(tenant_id,project_id,source_id,version,switched_at) VALUES(?,?,?,?,?)
        ON CONFLICT(tenant_id,project_id,source_id) DO UPDATE SET version=excluded.version, switched_at=excluded.switched_at
        """,
        (*_scope_values(scope), version, utc_now()),
    )


def _finish_batch(connection: sqlite3.Connection, batch_id: str, status: str, version: int | None, errors: list[dict[str, str]]) -> None:
    connection.execute(
        "UPDATE sync_batches SET status=?, completed_at=?, published_version=?, errors_json=? WHERE batch_id=?",
        (status, utc_now(), version, json.dumps(errors, ensure_ascii=False), batch_id),
    )


def _result(scope: SyncScope, batch_id: str, status: str, manifest_hash: str, version: int | None, records: list[dict[str, Any]], errors: list[dict[str, str]], database: Path) -> dict[str, Any]:
    return {
        "status": status,
        "batch_id": batch_id,
        **scope.__dict__,
        "manifest_hash": manifest_hash,
        "published_version": version,
        "candidate_record_count": len(records),
        "error_count": len(errors),
        "errors": errors,
        "database_path": str(database),
    }


def _scope_values(scope: SyncScope) -> tuple[str, str, str]:
    return scope.tenant_id, scope.project_id, scope.source_id


def _duplicate_record_keys(records: list[dict[str, Any]]) -> list[str]:
    seen: set[str] = set()
    duplicates: set[str] = set()
    for record in records:
        key = str(record["record_key"])
        if key in seen:
            duplicates.add(key)
        seen.add(key)
    return sorted(duplicates)


def _hash_json(value: Any) -> str:
    encoded = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()
