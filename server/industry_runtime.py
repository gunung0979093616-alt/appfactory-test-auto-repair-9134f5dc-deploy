from __future__ import annotations

import csv
import io
import json
import secrets
from datetime import datetime, timezone
from pathlib import Path
from typing import Any


MANIFEST = {'schema_version': '1.0', 'domain_id': 'automotive', 'runtime_variant': 'automotive', 'display_name': 'auto-repair', 'entity': 'record', 'search_tool': 'notification_delivery_status', 'detail_tool': 'notification_delivery_status', 'required_columns': ['id', 'name', 'status', 'updated_at'], 'tools': [{'name': 'schedule_operation', 'description': '確認資源、時間與影響範圍後建立排程。', 'action_kind': 'write', 'annotations': {'readOnlyHint': False, 'destructiveHint': False, 'idempotentHint': True, 'openWorldHint': False}, 'inputSchema': {'type': 'object', 'properties': {'query': {'type': 'string'}, 'page': {'type': 'integer', 'minimum': 1}, 'page_size': {'type': 'integer', 'minimum': 1, 'maximum': 100}, 'item_id': {'type': 'string'}, 'quantity': {'type': 'integer', 'minimum': 1}, 'confirmed': {'type': 'boolean'}, 'idempotency_key': {'type': 'string'}}, 'required': ['confirmed', 'idempotency_key'], 'additionalProperties': True}}, {'name': 'subscribe_alerts', 'description': '確認通知對象與管道後建立異常通知。', 'action_kind': 'write', 'annotations': {'readOnlyHint': False, 'destructiveHint': False, 'idempotentHint': True, 'openWorldHint': False}, 'inputSchema': {'type': 'object', 'properties': {'query': {'type': 'string'}, 'page': {'type': 'integer', 'minimum': 1}, 'page_size': {'type': 'integer', 'minimum': 1, 'maximum': 100}, 'item_id': {'type': 'string'}, 'quantity': {'type': 'integer', 'minimum': 1}, 'confirmed': {'type': 'boolean'}, 'idempotency_key': {'type': 'string'}}, 'required': ['confirmed', 'idempotency_key'], 'additionalProperties': True}}, {'name': 'create_booking', 'description': '重新檢查名額與輸入並確認後建立預約。', 'action_kind': 'transaction', 'annotations': {'readOnlyHint': False, 'destructiveHint': False, 'idempotentHint': True, 'openWorldHint': False}, 'inputSchema': {'type': 'object', 'properties': {'query': {'type': 'string'}, 'page': {'type': 'integer', 'minimum': 1}, 'page_size': {'type': 'integer', 'minimum': 1, 'maximum': 100}, 'booking_id': {'type': 'string'}, 'item_id': {'type': 'string'}, 'quantity': {'type': 'integer', 'minimum': 1}, 'confirmed': {'type': 'boolean'}, 'idempotency_key': {'type': 'string'}}, 'required': ['confirmed', 'idempotency_key'], 'additionalProperties': True}}, {'name': 'approve_task', 'description': '核准任務：依輸入、前置條件與成功條件執行，不以產業名稱代替需求判斷。', 'action_kind': 'transaction', 'annotations': {'readOnlyHint': False, 'destructiveHint': False, 'idempotentHint': True, 'openWorldHint': False}, 'inputSchema': {'type': 'object', 'properties': {'query': {'type': 'string'}, 'page': {'type': 'integer', 'minimum': 1}, 'page_size': {'type': 'integer', 'minimum': 1, 'maximum': 100}, 'item_id': {'type': 'string'}, 'quantity': {'type': 'integer', 'minimum': 1}, 'confirmed': {'type': 'boolean'}, 'idempotency_key': {'type': 'string'}}, 'required': ['confirmed', 'idempotency_key'], 'additionalProperties': True}}, {'name': 'send_notification', 'description': '發送通知：依輸入、前置條件與成功條件執行，不以產業名稱代替需求判斷。', 'action_kind': 'write', 'annotations': {'readOnlyHint': False, 'destructiveHint': False, 'idempotentHint': True, 'openWorldHint': False}, 'inputSchema': {'type': 'object', 'properties': {'query': {'type': 'string'}, 'page': {'type': 'integer', 'minimum': 1}, 'page_size': {'type': 'integer', 'minimum': 1, 'maximum': 100}, 'item_id': {'type': 'string'}, 'quantity': {'type': 'integer', 'minimum': 1}, 'confirmed': {'type': 'boolean'}, 'idempotency_key': {'type': 'string'}}, 'required': ['confirmed', 'idempotency_key'], 'additionalProperties': True}}, {'name': 'assign_technician', 'description': '指派技師：依輸入、前置條件與成功條件執行，不以產業名稱代替需求判斷。', 'action_kind': 'write', 'annotations': {'readOnlyHint': False, 'destructiveHint': False, 'idempotentHint': True, 'openWorldHint': False}, 'inputSchema': {'type': 'object', 'properties': {'query': {'type': 'string'}, 'page': {'type': 'integer', 'minimum': 1}, 'page_size': {'type': 'integer', 'minimum': 1, 'maximum': 100}, 'item_id': {'type': 'string'}, 'quantity': {'type': 'integer', 'minimum': 1}, 'confirmed': {'type': 'boolean'}, 'idempotency_key': {'type': 'string'}}, 'required': ['confirmed', 'idempotency_key'], 'additionalProperties': True}}, {'name': 'record_parts_usage', 'description': '零件使用紀錄：依輸入、前置條件與成功條件執行，不以產業名稱代替需求判斷。', 'action_kind': 'write', 'annotations': {'readOnlyHint': False, 'destructiveHint': False, 'idempotentHint': True, 'openWorldHint': False}, 'inputSchema': {'type': 'object', 'properties': {'query': {'type': 'string'}, 'page': {'type': 'integer', 'minimum': 1}, 'page_size': {'type': 'integer', 'minimum': 1, 'maximum': 100}, 'base_id': {'type': 'string'}, 'table_id': {'type': 'string'}, 'record_id': {'type': 'string'}, 'view_id': {'type': 'string'}, 'revision': {'type': 'string'}, 'cursor': {'type': 'string'}, 'field_values': {'type': 'object'}, 'records': {'type': 'array', 'items': {'type': 'object'}}, 'item_id': {'type': 'string'}, 'quantity': {'type': 'integer', 'minimum': 1}, 'confirmed': {'type': 'boolean'}, 'idempotency_key': {'type': 'string'}}, 'required': ['confirmed', 'idempotency_key'], 'additionalProperties': True}}, {'name': 'book_notification', 'description': 'book notification：依輸入、前置條件與成功條件執行，不以產業名稱代替需求判斷。', 'action_kind': 'write', 'annotations': {'readOnlyHint': False, 'destructiveHint': False, 'idempotentHint': True, 'openWorldHint': False}, 'inputSchema': {'type': 'object', 'properties': {'query': {'type': 'string'}, 'page': {'type': 'integer', 'minimum': 1}, 'page_size': {'type': 'integer', 'minimum': 1, 'maximum': 100}, 'item_id': {'type': 'string'}, 'quantity': {'type': 'integer', 'minimum': 1}, 'confirmed': {'type': 'boolean'}, 'idempotency_key': {'type': 'string'}}, 'required': ['confirmed', 'idempotency_key'], 'additionalProperties': True}}, {'name': 'cancel_notification', 'description': '取消通知：依輸入、前置條件與成功條件執行，不以產業名稱代替需求判斷。', 'action_kind': 'write', 'annotations': {'readOnlyHint': False, 'destructiveHint': False, 'idempotentHint': True, 'openWorldHint': False}, 'inputSchema': {'type': 'object', 'properties': {'query': {'type': 'string'}, 'page': {'type': 'integer', 'minimum': 1}, 'page_size': {'type': 'integer', 'minimum': 1, 'maximum': 100}, 'item_id': {'type': 'string'}, 'quantity': {'type': 'integer', 'minimum': 1}, 'confirmed': {'type': 'boolean'}, 'idempotency_key': {'type': 'string'}}, 'required': ['confirmed', 'idempotency_key'], 'additionalProperties': True}}, {'name': 'notification_delivery_status', 'description': '通知狀態：依輸入、前置條件與成功條件執行，不以產業名稱代替需求判斷。', 'action_kind': 'read', 'annotations': {'readOnlyHint': True, 'destructiveHint': False, 'idempotentHint': True, 'openWorldHint': False}, 'inputSchema': {'type': 'object', 'properties': {'query': {'type': 'string'}, 'page': {'type': 'integer', 'minimum': 1}, 'page_size': {'type': 'integer', 'minimum': 1, 'maximum': 100}, 'store_id': {'type': 'string'}, 'buyer_id': {'type': 'string'}, 'cart_id': {'type': 'string'}, 'variant_id': {'type': 'string'}, 'order_id': {'type': 'string'}}, 'required': [], 'additionalProperties': True}}, {'name': 'retry_notification', 'description': '重送通知：依輸入、前置條件與成功條件執行，不以產業名稱代替需求判斷。', 'action_kind': 'write', 'annotations': {'readOnlyHint': False, 'destructiveHint': False, 'idempotentHint': True, 'openWorldHint': False}, 'inputSchema': {'type': 'object', 'properties': {'query': {'type': 'string'}, 'page': {'type': 'integer', 'minimum': 1}, 'page_size': {'type': 'integer', 'minimum': 1, 'maximum': 100}, 'item_id': {'type': 'string'}, 'quantity': {'type': 'integer', 'minimum': 1}, 'confirmed': {'type': 'boolean'}, 'idempotency_key': {'type': 'string'}}, 'required': ['confirmed', 'idempotency_key'], 'additionalProperties': True}}, {'name': 'schedule_job', 'description': '排程工作：依輸入、前置條件與成功條件執行，不以產業名稱代替需求判斷。', 'action_kind': 'write', 'annotations': {'readOnlyHint': False, 'destructiveHint': False, 'idempotentHint': True, 'openWorldHint': False}, 'inputSchema': {'type': 'object', 'properties': {'query': {'type': 'string'}, 'page': {'type': 'integer', 'minimum': 1}, 'page_size': {'type': 'integer', 'minimum': 1, 'maximum': 100}, 'item_id': {'type': 'string'}, 'quantity': {'type': 'integer', 'minimum': 1}, 'confirmed': {'type': 'boolean'}, 'idempotency_key': {'type': 'string'}}, 'required': ['confirmed', 'idempotency_key'], 'additionalProperties': True}}, {'name': 'schedule_reminder', 'description': '排程提醒：依輸入、前置條件與成功條件執行，不以產業名稱代替需求判斷。', 'action_kind': 'write', 'annotations': {'readOnlyHint': False, 'destructiveHint': False, 'idempotentHint': True, 'openWorldHint': False}, 'inputSchema': {'type': 'object', 'properties': {'query': {'type': 'string'}, 'page': {'type': 'integer', 'minimum': 1}, 'page_size': {'type': 'integer', 'minimum': 1, 'maximum': 100}, 'item_id': {'type': 'string'}, 'quantity': {'type': 'integer', 'minimum': 1}, 'confirmed': {'type': 'boolean'}, 'idempotency_key': {'type': 'string'}}, 'required': ['confirmed', 'idempotency_key'], 'additionalProperties': True}}], 'capabilities': ['catalog_search_filter_pagination', 'detail_media_map_and_source', 'confirmed_user_actions', 'supplier_candidate_validate_publish', 'variants_slots_price_and_inventory', 'sandbox_checkout_or_booking', 'voucher_or_share_qr_when_applicable', 'typed_records_views_and_incremental_sync', 'workflow_state_activity_and_assignment', 'shipment_milestones_eta_and_exception_tracking'], 'oauth_scopes_by_tool': {}, 'publication_policy': 'candidate rows must validate before an explicit admin publish; never overwrite published data silently', 'truth_boundary': 'Generated local sandbox runtime and behavior tests only. Supplier publication, inventory, booking, checkout, voucher redemption, external media, OAuth, payment providers, deployment, and ChatGPT native rendering need their own authorization and live evidence.'}


class IndustryRuntime:
    def __init__(self, storage_root: Path):
        self.root = Path(storage_root)
        self.root.mkdir(parents=True, exist_ok=True)
        self.catalog_path = self.root / "runtime-catalog.json"
        self.candidates_path = self.root / "runtime-import-candidates.json"
        self.state_path = self.root / "runtime-state.json"

    def _read(self, path: Path, default: Any) -> Any:
        try:
            return json.loads(path.read_text(encoding="utf-8"))
        except (FileNotFoundError, json.JSONDecodeError, OSError):
            return default

    def _write(self, path: Path, payload: Any) -> None:
        path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    def catalog(self) -> dict[str, Any]:
        return self._read(self.catalog_path, {"version": "empty", "items": []})

    def search(self, arguments: dict[str, Any]) -> dict[str, Any]:
        query = str(arguments.get("query", "")).strip().casefold()
        filters = {key: value for key, value in arguments.items() if key not in {"query", "page", "page_size", "sort_by"} and value not in (None, "", [])}
        items = [item for item in self.catalog().get("items", []) if item.get("status") == "published"]
        if query:
            items = [item for item in items if query in json.dumps(item, ensure_ascii=False).casefold()]
        for key, value in filters.items():
            if MANIFEST.get("runtime_variant") == "travel" and key == "date":
                items = [item for item in items if value in item.get("available_dates", [])]
            elif MANIFEST.get("runtime_variant") == "travel" and key == "min_rating":
                items = [item for item in items if float(item.get("rating") or 0) >= float(value)]
            elif MANIFEST.get("runtime_variant") == "travel" and key == "max_distance_km":
                items = [item for item in items if item.get("distance_km") is not None and float(item["distance_km"]) <= float(value)]
            elif MANIFEST.get("runtime_variant") == "travel" and key == "activity_type":
                items = [item for item in items if str(value).casefold() in str(item.get("category", "")).casefold()]
            elif MANIFEST.get("runtime_variant") == "travel" and key == "instant_confirmation":
                items = [item for item in items if item.get("instant_confirmation") is value]
            elif key in {"max_price", "min_price"}:
                boundary = float(value)
                items = [item for item in items if (float(item.get("price", 0)) <= boundary if key == "max_price" else float(item.get("price", 0)) >= boundary)]
            elif key in {"confirmed", "quantity", "item_id", "authorized_scopes", "idempotency_key", "actor_id", "actor_role", "notification_consent"}:
                continue
            else:
                items = [item for item in items if str(value).casefold() in str(item.get(key, "")).casefold()]
        if MANIFEST.get("runtime_variant") == "travel":
            sort_by = str(arguments.get("sort_by") or "relevance")
            sort_keys = {
                "price_asc": (lambda item: float(item.get("price") or 0), False),
                "price_desc": (lambda item: float(item.get("price") or 0), True),
                "rating_desc": (lambda item: float(item.get("rating") or 0), True),
                "distance_asc": (lambda item: float(item.get("distance_km") if item.get("distance_km") is not None else float("inf")), False),
            }
            if sort_by in sort_keys:
                key_fn, reverse = sort_keys[sort_by]
                items.sort(key=key_fn, reverse=reverse)
        page = max(1, int(arguments.get("page", 1)))
        page_size = min(100, max(1, int(arguments.get("page_size", 20))))
        start = (page - 1) * page_size
        selected = items[start:start + page_size]
        return self._ui("success" if selected else "empty", "搜尋結果", selected, {"page": page, "page_size": page_size, "total_count": len(items), "next_cursor": str(page + 1) if start + page_size < len(items) else None})

    def get(self, item_id: str) -> dict[str, Any]:
        item = next((row for row in self.catalog().get("items", []) if item_id in {str(row.get("id") or ""), str(row.get("record_id") or ""), str(row.get("card_id") or ""), str(row.get("shipment_id") or ""), str(row.get("tracking_number") or "")} and row.get("status") == "published"), None)
        if not item:
            return self._ui("empty", "找不到項目", [], {"page": 1, "page_size": 1, "total_count": 0, "next_cursor": None})
        return self._ui("success", str(item.get("name") or "項目詳情"), [item], {"page": 1, "page_size": 1, "total_count": 1, "next_cursor": None}, content_type="detail")

    def _call_automotive_operations(self, tool_name: str, arguments: dict[str, Any], state: dict[str, Any]) -> dict[str, Any]:
        state.setdefault("vehicle_intakes", [])
        state.setdefault("repair_orders", [])
        state.setdefault("idempotency", {})
        state.setdefault("activity", [])
        read_tools = {"diagnose_wait_bottlenecks", "list_repair_work_queue", "get_repair_order", "match_technician_skills"}
        if tool_name == "diagnose_wait_bottlenecks":
            shop_id = str(arguments.get("shop_id") or "")
            active = [row for row in state["repair_orders"] if row.get("shop_id") == shop_id and row.get("status") != "completed"]
            evidence = [
                {"stage": "vehicle_intake", "median_wait_minutes": 18, "open_items": len(state["vehicle_intakes"])},
                {"stage": "quote_approval", "median_wait_minutes": 42, "open_items": sum(row.get("status") == "waiting_quote_approval" for row in active)},
                {"stage": "parts_wait", "median_wait_minutes": 95, "open_items": sum(row.get("status") == "waiting_parts" for row in active)},
                {"stage": "repair_bay", "median_wait_minutes": 27, "open_items": sum(row.get("status") == "repairing" for row in active)},
            ]
            ranked = sorted(evidence, key=lambda row: row["median_wait_minutes"], reverse=True)
            return {"status": "ok", "shop_id": shop_id, "root_cause_verified": False, "evidence": evidence, "ranked_hypotheses": ranked, "next_step": "collect_real_event_timestamps_before_freezing_solution"}
        if tool_name == "list_repair_work_queue":
            shop_id = str(arguments.get("shop_id") or "")
            requested_status = str(arguments.get("status") or "")
            rows = [row for row in state["repair_orders"] if row.get("shop_id") == shop_id and (not requested_status or row.get("status") == requested_status)]
            return {"status": "ok", "shop_id": shop_id, "items": rows, "total_count": len(rows)}
        repair_order_id = str(arguments.get("repair_order_id") or "")
        order = next((row for row in state["repair_orders"] if row.get("id") == repair_order_id), None)
        if tool_name == "get_repair_order":
            return {"status": "ok", "repair_order": order} if order else {"status": "not_found", "repair_order_id": repair_order_id}
        if tool_name == "match_technician_skills":
            if not order:
                return {"status": "not_found", "repair_order_id": repair_order_id}
            return {"status": "ok", "repair_order_id": repair_order_id, "candidates": [{"technician_id": "tech-demo-01", "matched_skills": order.get("inspection_items", []), "available": True}], "assignment_changed": False}
        if tool_name not in read_tools and arguments.get("confirmed") is not True:
            return {"status": "confirmation_required", "tool_name": tool_name, "external_write_completed": False}
        idempotency_key = str(arguments.get("idempotency_key") or "")
        if not idempotency_key:
            return {"status": "idempotency_key_required", "tool_name": tool_name, "external_write_completed": False}
        if idempotency_key in state["idempotency"]:
            return state["idempotency"][idempotency_key]
        if tool_name == "create_vehicle_intake":
            intake = {
                "id": "intake_" + secrets.token_hex(6), "shop_id": str(arguments.get("shop_id") or ""),
                "plate_number": str(arguments.get("plate_number") or ""), "mileage_km": int(arguments.get("mileage_km") or 0),
                "customer_concern": str(arguments.get("customer_concern") or ""), "status": "intake_recorded",
            }
            state["vehicle_intakes"].append(intake)
            response = {"status": "created", "vehicle_intake": intake}
        elif tool_name == "create_repair_order":
            intake = next((row for row in state["vehicle_intakes"] if row.get("id") == str(arguments.get("intake_id") or "")), None)
            if not intake:
                return {"status": "intake_not_found", "intake_id": arguments.get("intake_id")}
            order = {"id": "repair_" + secrets.token_hex(6), "intake_id": intake["id"], "shop_id": intake["shop_id"], "plate_number": intake["plate_number"], "inspection_items": list(arguments.get("inspection_items") or []), "status": "inspection", "timeline": []}
            state["repair_orders"].append(order)
            response = {"status": "created", "repair_order": order}
        else:
            if not order:
                return {"status": "not_found", "repair_order_id": repair_order_id}
            if tool_name == "assign_repair_bay":
                bay_id = str(arguments.get("bay_id") or "")
                occupied = any(row.get("bay_id") == bay_id and row.get("status") not in {"ready_for_handover", "completed"} and row.get("id") != repair_order_id for row in state["repair_orders"])
                if occupied:
                    return {"status": "bay_occupied", "bay_id": bay_id}
                order["bay_id"] = bay_id
            elif tool_name == "mark_parts_wait":
                order["status"] = "waiting_parts"
                order.setdefault("parts_wait", []).append({"part_name": arguments.get("part_name"), "estimated_arrival": arguments.get("estimated_arrival")})
            elif tool_name == "request_quote_approval":
                order["status"] = "waiting_quote_approval"
                order["quote"] = {"amount": arguments.get("amount"), "items": list(arguments.get("quote_items") or []), "approval": "pending"}
            elif tool_name == "update_repair_status":
                order["status"] = str(arguments.get("status") or order.get("status"))
            elif tool_name == "notify_repair_customer":
                notification = {"channel": arguments.get("channel"), "message": arguments.get("message"), "delivery_status": "sandbox_recorded", "external_delivery_completed": False}
                order.setdefault("notifications", []).append(notification)
            elif tool_name == "complete_vehicle_handover":
                if arguments.get("quality_check_passed") is not True:
                    return {"status": "quality_check_required", "repair_order_id": repair_order_id}
                order["status"] = "completed"
            order.setdefault("timeline", []).append({"action": tool_name, "status": order.get("status")})
            response = {"status": "ok", "tool_name": tool_name, "repair_order": order}
        state["idempotency"][idempotency_key] = response
        state["activity"].append({"tool_name": tool_name, "idempotency_key": idempotency_key})
        self._write(self.state_path, state)
        return response

    def call(self, tool_name: str, arguments: dict[str, Any]) -> dict[str, Any]:
        state = self._read(self.state_path, {"saved": [], "cart": [], "bookings": [], "vouchers": {}, "records": [], "activity": [], "subscriptions": [], "jobs": {}, "sessions": {}, "drafts": [], "designs": {}, "assets": {}, "comments": [], "idempotency": {}})
        for key, default in [("saved", []), ("cart", []), ("bookings", []), ("vouchers", {}), ("records", []), ("activity", []), ("subscriptions", []), ("jobs", {}), ("sessions", {}), ("drafts", []), ("designs", {}), ("assets", {}), ("comments", []), ("idempotency", {})]:
            state.setdefault(key, default)
        automotive_tools = {item["name"] for item in json.loads(Path("generated_mcp/automotive-operations-tools.json").read_text(encoding="utf-8"))["tools"]}
        if MANIFEST.get("runtime_variant") == "automotive" and tool_name in automotive_tools:
            return self._call_automotive_operations(tool_name, arguments, state)
        if MANIFEST.get("runtime_variant") in ('abandoned_cart_sms_recovery', 'ai_customer_support', 'ai_data_analysis_workspace', 'ai_interview_training', 'arborist_certification_learning', 'daily_task_checkin', 'ecommerce_product_copy', 'email_attachment_print_automation', 'email_outreach_crm', 'llm_api_gateway', 'long_video_content_repurposing', 'marketplace_ranking_optimization', 'merchant_payment_links', 'mobile_app_market_intelligence', 'openclaw_deployment_manager', 'prompt_optimization_lab', 'reddit_purchase_intent', 'research_focus_group_recruitment', 'resume_job_matching', 'sales_lead_research', 'security_compliance_management', 'short_link_conversion_analytics', 'small_business_ordering_reservation', 'social_media_scheduling', 'tiktok_comment_intelligence', 'toeic_exam_training', 'video_translation_dubbing', 'website_translation_operations'):
            action_kind = next((item["action_kind"] for item in MANIFEST["tools"] if item["name"] == tool_name), "read")
            if action_kind in {"write", "transaction"} and arguments.get("confirmed") is not True:
                return {"status": "confirmation_required", "tool_name": tool_name, "external_write_completed": False}
            idempotency_key = str(arguments.get("idempotency_key") or "")
            if action_kind in {"write", "transaction"} and not idempotency_key:
                return {"status": "idempotency_key_required", "tool_name": tool_name, "external_write_completed": False}
            if idempotency_key and idempotency_key in state["idempotency"]:
                return state["idempotency"][idempotency_key]
            if tool_name == MANIFEST.get("search_tool"):
                response = self.search(arguments)
                response["capability"] = tool_name
                response["source_scope"] = "current_project_authorized_sources_only"
                return response
            read_results = {
                "find_company_contacts": {"contacts": [], "source_verified": False},
                "get_company_sources": {"sources": [], "replaceable": True},
                "verify_contact_details": {"verification_status": "needs_provider_check", "verified_at": None},
                "detect_stale_contacts": {"stale_contact_ids": [], "rule": "source_age_and_verification_status"},
                "deduplicate_leads": {"duplicate_groups": [], "match_keys": ["company_domain", "contact_email"]},
                "filter_leads": {"lead_ids": [], "filters_applied": arguments.get("filters", {})},
                "get_lead_history": {"events": state.get("activity", [])},
                "get_data_quality_status": {"quality_status": "source_rows_required", "dimensions": ["completeness", "freshness", "source_traceability"]},
                "extract_page_content": {"content_blocks": [], "source_content_hash": arguments.get("source_content_hash")},
                "preview_translation": {"preview_status": "draft", "translation_id": arguments.get("translation_id"), "published": False},
                "evaluate_mail_rules": {"matched": False, "rule_id": arguments.get("rule_id"), "rule_version": arguments.get("rule_version")},
                "classify_attachments": {"classification": "needs_scan", "attachment_id": arguments.get("attachment_id"), "safe_to_print": False},
                "deduplicate_mail_jobs": {"duplicate": False, "dedupe_key": arguments.get("attachment_hash")},
                "get_mail_processing_audit": {"events": state.get("activity", []), "project_scoped": True},
            }
            if action_kind == "read":
                return {"status": "ok", "tool_name": tool_name, **read_results.get(tool_name, {"items": []})}
            status_by_tool = {
                "normalize_company_records": "normalization_prepared",
                "export_leads": "export_prepared",
                "replace_company_data_source": "source_replacement_prepared",
                "translate_content": "translation_draft_prepared",
                "preserve_seo_metadata": "seo_mapping_prepared",
                "edit_translation": "translation_revision_prepared",
                "publish_translation": "publication_prepared",
                "sync_changed_pages": "change_sync_checkpoint_prepared",
                "configure_mail_rules": "mail_rules_prepared",
                "enqueue_print_job": "print_job_queued",
                "retry_failed_print_job": "print_retry_queued",
                "rerun_mail_job": "mail_job_rerun_queued",
            }
            response = {
                "status": "ok",
                "operation_status": status_by_tool.get(tool_name, "sandbox_write_prepared"),
                "tool_name": tool_name,
                "sandbox": True,
                "provider_authorized": False,
                "external_write_completed": False,
                "requirement_variant": MANIFEST.get("runtime_variant"),
            }
            state["idempotency"][idempotency_key] = response
            state["activity"].append({"tool_name": tool_name, "operation_status": response["operation_status"]})
            self._write(self.state_path, state)
            return response
        if MANIFEST.get("runtime_variant") == "google_workspace_automation":
            action_kind = next((item["action_kind"] for item in MANIFEST["tools"] if item["name"] == tool_name), "read")
            if action_kind in {"write", "transaction"} and not arguments.get("confirmed"):
                return {"status": "confirmation_required", "tool_name": tool_name, "message": "請先確認帳號、目標、內容與外部影響；目前尚未寫入。"}
            required_scopes = set(MANIFEST.get("oauth_scopes_by_tool", {}).get(tool_name, []))
            authorized_scopes = set(arguments.get("authorized_scopes") or [])
            if not required_scopes.issubset(authorized_scopes):
                return {"status": "authorization_required", "tool_name": tool_name, "missing_scopes": sorted(required_scopes - authorized_scopes), "credentials_requested_in_chat": False}
            if action_kind in {"write", "transaction"}:
                idempotency_key = str(arguments.get("idempotency_key") or "")
                if not idempotency_key:
                    return {"status": "idempotency_key_required", "message": "外部寫入需要唯一操作碼。"}
                if idempotency_key in state["idempotency"]:
                    return state["idempotency"][idempotency_key]
                status_by_tool = {
                    "append_sheet": "sandbox_rows_prepared", "update_sheet": "sandbox_update_prepared",
                    "create_document": "sandbox_document_prepared", "create_pdf": "sandbox_pdf_prepared",
                    "save_to_drive": "sandbox_drive_save_prepared", "send_email": "sandbox_email_prepared",
                    "create_calendar_event": "sandbox_event_prepared", "generate_and_send_report": "sandbox_workflow_prepared",
                }
                response = {"status": status_by_tool.get(tool_name, "sandbox_write_prepared"), "tool_name": tool_name, "provider_authorized": False, "external_write_completed": False, "operation_order": ["create_document", "create_pdf", "save_to_drive", "send_email", "append_sheet", "create_calendar_event"] if tool_name == "generate_and_send_report" else [tool_name]}
                state["idempotency"][idempotency_key] = response
                state["activity"].append({"tool_name": tool_name, "status": response["status"]})
                self._write(self.state_path, state)
                return response
        if MANIFEST.get("runtime_variant") == "design_creation":
            if tool_name == "list_design_templates":
                return self.search(arguments)
            if tool_name == "get_design":
                design_id = str(arguments.get("design_id") or "")
                design = state["designs"].get(design_id)
                if not design:
                    design = next((row for row in self.catalog().get("items", []) if design_id in {str(row.get("id") or ""), str(row.get("design_id") or "")}), None)
                return design or {"status": "not_found", "design_id": design_id}
            if tool_name == "get_design_job":
                job_id = str(arguments.get("job_id") or "")
                return state["jobs"].get(job_id) or {"status": "not_found", "job_id": job_id}
            action_kind = next((item["action_kind"] for item in MANIFEST["tools"] if item["name"] == tool_name), "read")
            if action_kind in {"write", "transaction"} and arguments.get("confirmed") is not True:
                return {"status": "confirmation_required", "tool_name": tool_name, "external_write_completed": False}
            idempotency_key = str(arguments.get("idempotency_key") or "")
            if not idempotency_key:
                return {"status": "idempotency_key_required", "tool_name": tool_name}
            if idempotency_key in state["idempotency"]:
                return state["idempotency"][idempotency_key]
            if tool_name == "upload_design_asset" and arguments.get("rights_declaration") is not True:
                return {"status": "rights_declaration_required", "asset_saved": False}
            if tool_name == "apply_brand_kit" and arguments.get("brand_authorized") is not True:
                return {"status": "brand_authorization_required", "design_changed": False}
            if tool_name == "create_design":
                design_id = "design_" + secrets.token_hex(6)
                design = {"status": "draft", "design_id": design_id, "title": arguments.get("title"), "template_id": arguments.get("template_id"), "version_id": "v1", "locale": arguments.get("locale", "zh-Hant")}
                state["designs"][design_id] = design
                response = design
            elif tool_name == "upload_design_asset":
                asset_id = "asset_" + secrets.token_hex(6)
                response = {"status": "sandbox_asset_saved", "asset_id": asset_id, "public": False, "rights_recorded": True}
                state["assets"][asset_id] = response
            elif tool_name == "bulk_create_designs":
                job_id = "job_" + secrets.token_hex(6)
                response = {"status": "queued", "job_id": job_id, "item_count": len(arguments.get("data_rows") or []), "result_url": None}
                state["jobs"][job_id] = response
            elif tool_name == "export_design":
                job_id = "job_" + secrets.token_hex(6)
                response = {"status": "queued", "job_id": job_id, "design_id": arguments.get("design_id"), "format": arguments.get("format"), "result_url": None}
                state["jobs"][job_id] = response
            elif tool_name == "add_design_comment":
                response = {"status": "comment_added", "design_id": arguments.get("design_id"), "comment_id": "comment_" + secrets.token_hex(6)}
                state["comments"].append({**response, "comment": arguments.get("comment")})
            else:
                design_id = str(arguments.get("design_id") or "")
                design = state["designs"].get(design_id)
                if not design:
                    return {"status": "not_found", "design_id": design_id}
                if tool_name == "edit_design" and arguments.get("base_version_id") != design.get("version_id"):
                    return {"status": "version_conflict", "current_version_id": design.get("version_id")}
                version_number = int(str(design.get("version_id") or "v1").lstrip("v")) + 1
                design["version_id"] = f"v{version_number}"
                design["last_action"] = tool_name
                if tool_name == "resize_design":
                    design.update({"width": arguments.get("width"), "height": arguments.get("height")})
                if tool_name == "apply_brand_kit":
                    design["brand_kit_id"] = arguments.get("brand_kit_id")
                if tool_name == "restore_design_version":
                    design["restored_from_version_id"] = arguments.get("version_id")
                response = {"status": "updated", "design_id": design_id, "version_id": design["version_id"], "action": tool_name}
            state["idempotency"][idempotency_key] = response
            self._write(self.state_path, state)
            return response
        if tool_name == "search" and MANIFEST["domain_id"] == "structured_data":
            result = self.search(arguments)
            rows = result["structuredContent"]["items"]
            payload = {"results": [{"id": row["id"], "title": row["title"], "url": row["fields"].get("source_url") or f"https://example.com/records/{row['id']}"} for row in rows]}
            result["content"] = [{"type": "text", "text": json.dumps(payload, ensure_ascii=False)}]
            return result
        if tool_name == "fetch" and MANIFEST["domain_id"] == "structured_data":
            item_id = str(arguments.get("id") or arguments.get("record_id") or arguments.get("item_id") or "")
            available = self.catalog().get("items", []) + state.get("records", [])
            item = next((row for row in available if row.get("id") == item_id and row.get("status") in {"published", "draft"}), None)
            if not item:
                return {"status": "not_found", "content": [{"type": "text", "text": json.dumps({"id": item_id, "title": "找不到資料", "text": "", "url": ""}, ensure_ascii=False)}]}
            result = self.get(item_id)
            payload = {"id": item_id, "title": str(item.get("name") or item_id), "text": json.dumps(item, ensure_ascii=False), "url": item.get("source_url") or f"https://example.com/records/{item_id}", "metadata": {"updated_at": item.get("updated_at"), "revision": item.get("revision")}}
            result["content"] = [{"type": "text", "text": json.dumps(payload, ensure_ascii=False)}]
            return result
        if tool_name == "list_views":
            return {"status": "ok", "items": [{"id": "all", "name": "全部資料", "table_id": str(arguments.get("table_id") or "table-demo"), "read_only": True}], "total_count": 1}
        if tool_name == "list_shipment_events":
            item_id = str(arguments.get("shipment_id") or arguments.get("tracking_number") or arguments.get("item_id") or "")
            item = next((row for row in self.catalog().get("items", []) if item_id in {str(row.get("id") or ""), str(row.get("shipment_id") or ""), str(row.get("tracking_number") or "")} and row.get("status") == "published"), None)
            if not item:
                return {"status": "not_found", "item_id": item_id, "items": []}
            return {"status": "ok", "item_id": item_id, "items": item.get("milestones", []), "exceptions": item.get("exceptions", []), "last_event_at": item.get("last_event_at"), "event_source": item.get("event_source"), "eta_is_estimate": item.get("eta_is_estimate") is True}
        if tool_name == "list_board":
            board_id = str(arguments.get("board_id") or "")
            items = [row for row in self.catalog().get("items", []) if row.get("status") == "published" and (not board_id or row.get("board_id") == board_id)]
            return {"status": "ok", "board_id": board_id, "items": sorted(items, key=lambda row: (str(row.get("list_id") or ""), int(row.get("position") or 0))), "total_count": len(items)}
        if tool_name == "list_activity":
            return {"status": "ok", "items": state.get("activity", []), "total_count": len(state.get("activity", []))}
        if tool_name == "validate_quote_inputs":
            missing = [key for key in ["residence", "destinations", "travel_dates", "traveller_ages"] if arguments.get(key) in (None, "", [])]
            return {"status": "valid" if not missing else "needs_input", "missing": missing, "identity_required": False, "sensitive_health_data_requested": False}
        if tool_name == "search_policy_quotes":
            as_of = str(arguments.get("as_of") or datetime.now(timezone.utc).isoformat())
            items = [row for row in self.catalog().get("items", []) if row.get("status") == "published" and str(row.get("expires_at") or "") > as_of]
            result = self._ui("success" if items else "empty", "有效報價", items, {"page": 1, "page_size": len(items), "total_count": len(items), "next_cursor": None})
            result["comparison_scope"] = "current_project_authorized_sources_only"
            result["contract_truth"] = "official_policy_document"
            return result
        if tool_name == "start_live_offer_search":
            session_id = "session_" + secrets.token_hex(6)
            search_result = self.search(arguments)
            session = {"session_id": session_id, "status": "partial", "poll_count": 0, "query": arguments, "items": self.catalog().get("items", []), "provider_authorized": False, "structuredContent": search_result.get("structuredContent")}
            state["sessions"][session_id] = session
            self._write(self.state_path, state)
            return session
        if tool_name == "poll_live_offer_search":
            session_id = str(arguments.get("session_id") or "")
            session = state["sessions"].get(session_id)
            if not session:
                return {"status": "session_expired", "session_id": session_id, "restart_required": True}
            session["poll_count"] += 1
            session["status"] = "complete" if session["poll_count"] >= 2 else "partial"
            self._write(self.state_path, state)
            return session
        if tool_name in {"get_media_job", "get_analysis_job"}:
            job_id = str(arguments.get("job_id") or arguments.get("item_id") or "")
            return state["jobs"].get(job_id) or {"status": "not_found", "job_id": job_id}
        if tool_name == "get_learning_progress" and arguments.get("member_authorized") is not True:
            return {"status": "permission_denied", "message": "私人學習進度需要會員明確授權。"}
        if tool_name == "list_media_outputs":
            items = [job for job in state["jobs"].values() if job.get("durable_output_ready")]
            return self._ui("success" if items else "empty", "媒體輸出", items, {"page": 1, "page_size": len(items), "total_count": len(items), "next_cursor": None})
        if tool_name == "get_booking_link" and MANIFEST.get("runtime_variant") == "travel":
            item_id = str(arguments.get("item_id") or "")
            item = next((row for row in self.catalog().get("items", []) if str(row.get("id")) == item_id and row.get("status") == "published"), None)
            if not item:
                return {"status": "not_found", "item_id": item_id}
            url = str(item.get("booking_url") or "")
            if not url.startswith("https://") or item.get("booking_link_verified") is not True:
                return {"status": "not_available", "item_id": item_id, "booking_url": None, "reason": "verified_booking_link_required"}
            return {"status": "ok", "item_id": item_id, "booking_url": url, "external_booking_not_completed": True}
        if tool_name == MANIFEST.get("search_tool") or tool_name.startswith("list_"):
            return self.search(arguments)
        if (tool_name == MANIFEST.get("detail_tool") or tool_name.startswith(("get_", "track_"))) and tool_name not in {"get_cart", "get_voucher", "get_booking", "get_media_job", "get_analysis_job"}:
            return self.get(str(arguments.get("item_id") or arguments.get("quote_id") or arguments.get("offer_id") or arguments.get("event_id") or arguments.get("meeting_id") or arguments.get("shipment_id") or arguments.get("tracking_number") or arguments.get("id") or ""))
        if tool_name == "get_cart":
            return {"status": "ok", "items": state["cart"], "total_count": len(state["cart"])}
        if tool_name == "get_voucher":
            voucher_id = str(arguments.get("voucher_id") or arguments.get("item_id") or "")
            voucher = state["vouchers"].get(voucher_id)
            return voucher or {"status": "not_found", "voucher_id": voucher_id}
        if tool_name == "get_booking":
            booking_id = str(arguments.get("booking_id") or arguments.get("item_id") or "")
            return next((booking for booking in state["bookings"] if booking.get("id") == booking_id), {"status": "not_found", "booking_id": booking_id})
        if tool_name.startswith("share_"):
            item_id = str(arguments.get("item_id") or "")
            return {"status": "ok", "item_id": item_id, "qr_payload": f"https://example.com/items/{item_id}", "note": "前端可將 qr_payload 轉成 QR；只分享專案自有網址。"}
        if tool_name.startswith("redeem_"):
            return self.redeem(str(arguments.get("voucher_id") or arguments.get("item_id") or ""), confirmed=bool(arguments.get("confirmed")))
        action_kind = next((item["action_kind"] for item in MANIFEST["tools"] if item["name"] == tool_name), "read")
        if action_kind == "read":
            item_id = str(arguments.get("item_id") or arguments.get("quote_id") or arguments.get("offer_id") or arguments.get("event_id") or arguments.get("meeting_id") or arguments.get("id") or "")
            return self.get(item_id) if item_id else self.search(arguments)
        if action_kind in {"write", "transaction"} and not arguments.get("confirmed"):
            return {"status": "confirmation_required", "tool_name": tool_name, "message": "請先確認項目、數量、日期與影響範圍；目前尚未寫入。"}
        if MANIFEST["domain_id"] == "workflow" and action_kind == "write" and arguments.get("actor_role") not in {"owner_admin", "member"}:
            return {"status": "permission_denied", "message": "看板寫入需要已授權的成員身分。"}
        if tool_name == "acknowledge_shipment_exception" and arguments.get("actor_role") not in {"owner_admin", "operator"}:
            return {"status": "permission_denied", "message": "只有已授權管理者或物流人員可以確認異常。"}
        if tool_name == "subscribe_shipment_updates" and arguments.get("notification_consent") is not True:
            return {"status": "notification_consent_required", "message": "建立通知前需要使用者明確同意。"}
        if tool_name == "submit_media_job" and arguments.get("rights_declaration") is not True:
            return {"status": "rights_declaration_required", "message": "送出生成前需確認素材權利；涉及真人時另需肖像同意依據。"}
        if tool_name == "submit_media_job" and arguments.get("confirmed_cost") is not True:
            return {"status": "cost_confirmation_required", "message": "付費生成前需確認成本與保留額度。"}
        if tool_name == "create_event_draft" and not str(arguments.get("timezone") or "").strip():
            return {"status": "timezone_required", "message": "建立行程前需要明確時區。"}
        if tool_name == "cancel_booking":
            booking_id = str(arguments.get("booking_id") or arguments.get("item_id") or "")
            booking = next((row for row in state["bookings"] if row.get("id") == booking_id), None)
            if not booking:
                return {"status": "not_found", "booking_id": booking_id}
            if booking.get("status") == "cancelled":
                return {"status": "already_cancelled", "booking_id": booking_id}
            booking["status"] = "cancelled"
            self._write(self.state_path, state)
            return {"status": "cancelled", "booking_id": booking_id}
        idempotency_key = str(arguments.get("idempotency_key") or "")
        if action_kind in {"write", "transaction"} and not idempotency_key:
            return {"status": "idempotency_key_required", "message": "寫入操作需要唯一操作碼，避免重複建立或更新。"}
        if idempotency_key and idempotency_key in state["idempotency"]:
            return state["idempotency"][idempotency_key]
        item_id = str(arguments.get("item_id") or arguments.get("record_id") or arguments.get("card_id") or arguments.get("shipment_id") or arguments.get("tracking_number") or arguments.get("id") or "")
        if tool_name in {"submit_media_job", "run_analysis_job"}:
            job_id = "job_" + secrets.token_hex(6)
            response = {"status": "queued", "job_id": job_id, "job_type": arguments.get("job_type") or ("analysis" if tool_name == "run_analysis_job" else "media"), "provider_authorized": False, "reserved_quota": True, "committed_quota": False, "durable_output_ready": False, "retryable": True}
            state["jobs"][job_id] = response
            state["idempotency"][idempotency_key] = response
            self._write(self.state_path, state)
            return response
        if tool_name in {"cancel_media_job", "retry_media_job"}:
            job_id = str(arguments.get("job_id") or arguments.get("item_id") or "")
            job = state["jobs"].get(job_id)
            if not job:
                return {"status": "not_found", "job_id": job_id}
            job["status"] = "cancelled" if tool_name == "cancel_media_job" else "queued"
            job["quota_compensation_pending"] = tool_name == "cancel_media_job"
            response = dict(job)
            state["idempotency"][idempotency_key] = response
            self._write(self.state_path, state)
            return response
        if tool_name == "create_event_draft":
            draft = {"id": "event_draft_" + secrets.token_hex(6), "status": "draft", "timezone": arguments.get("timezone"), "start_at": arguments.get("start_at"), "end_at": arguments.get("end_at"), "attendees": arguments.get("attendees", []), "invitation_sent": False}
            state["drafts"].append(draft); state["idempotency"][idempotency_key] = draft; self._write(self.state_path, state); return draft
        if tool_name == "create_cart":
            response = {"status": "cart_created", "cart_id": "cart_" + secrets.token_hex(6), "store_id": arguments.get("store_id"), "buyer_id": arguments.get("buyer_id"), "payment_status": "not_started"}
            state["idempotency"][idempotency_key] = response; self._write(self.state_path, state); return response
        if tool_name in {"prepare_provider_handoff", "prepare_checkout_redirect", "confirm_enrollment_handoff"}:
            response = {"status": "sandbox_redirect_ready", "url": "https://example.com/provider-handoff", "allowlisted_https": True, "booking_completed": False, "payment_completed": False, "provider_authorized": False}
            state["idempotency"][idempotency_key] = response; self._write(self.state_path, state); return response
        if tool_name in {"create_record", "create_card"}:
            if tool_name == "create_record" and not isinstance(arguments.get("field_values", {}), dict):
                return {"status": "field_validation_failed", "errors": [{"field": "field_values", "message": "必須是物件格式"}]}
            created_id = str(arguments.get("record_id") or arguments.get("card_id") or arguments.get("id") or ("record_" + secrets.token_hex(6)))
            created = {"id": created_id, "name": str(arguments.get("title") or arguments.get("name") or "未命名項目"), "status": "draft", "board_id": arguments.get("board_id"), "list_id": arguments.get("list_id"), "field_values": arguments.get("field_values", {}), "revision": "1"}
            state["records"].append(created)
            response = {"status": "created", "tool_name": tool_name, "item": created}
            state["idempotency"][idempotency_key] = response
            self._write(self.state_path, state)
            return response
        if tool_name == "batch_upsert_records":
            rows = arguments.get("records") if isinstance(arguments.get("records"), list) else []
            accepted = [row for row in rows if isinstance(row, dict) and row.get("id") and isinstance(row.get("field_values", {}), dict)]
            failed = [{"index": index, "message": "id 與 field_values 格式不完整"} for index, row in enumerate(rows) if row not in accepted]
            state["records"].extend({"id": str(row["id"]), "name": str(row.get("name") or row["id"]), "status": "draft", "field_values": row.get("field_values", {}), "revision": "1"} for row in accepted)
            response = {"status": "accepted" if not failed else "partial_failure", "tool_name": tool_name, "accepted_count": len(accepted), "failed_count": len(failed), "errors": failed, "sandbox": True}
            state["idempotency"][idempotency_key] = response
            self._write(self.state_path, state)
            return response
        if tool_name == "sync_changes":
            response = {"status": "checkpoint_saved", "next_cursor": str(arguments.get("cursor") or "cursor-demo-next"), "provider_authorized": False, "live_webhook_received": False}
            state["idempotency"][idempotency_key] = response
            self._write(self.state_path, state)
            return response
        if tool_name == "update_record":
            catalog = self.catalog()
            candidates = catalog.get("items", []) + state.get("records", [])
            record = next((row for row in candidates if item_id in {str(row.get("id") or ""), str(row.get("record_id") or "")}), None)
            if not record:
                return {"status": "not_found", "item_id": item_id}
            if str(arguments.get("revision") or "") != str(record.get("revision") or "1"):
                return {"status": "revision_conflict", "current_revision": str(record.get("revision") or "1")}
            if not isinstance(arguments.get("field_values", {}), dict):
                return {"status": "field_validation_failed", "errors": [{"field": "field_values", "message": "必須是物件格式"}]}
            record["field_values"] = {**record.get("field_values", {}), **arguments.get("field_values", {})}
            record["revision"] = str(int(record.get("revision") or 1) + 1)
            response = {"status": "ok", "tool_name": tool_name, "item_id": item_id, "revision": record["revision"]}
            state["idempotency"][idempotency_key] = response
            self._write(self.state_path, state)
            self._write(self.catalog_path, catalog)
            return response
        item = next((row for row in self.catalog().get("items", []) if item_id in {str(row.get("id") or ""), str(row.get("record_id") or ""), str(row.get("card_id") or ""), str(row.get("shipment_id") or ""), str(row.get("tracking_number") or "")}), None)
        if not item:
            return {"status": "not_found", "item_id": item_id}
        quantity = max(1, int(arguments.get("quantity", 1)))
        if action_kind == "transaction" and int(item.get("inventory", 0)) < quantity:
            return {"status": "inventory_unavailable", "available": int(item.get("inventory", 0))}
        if tool_name == "add_to_cart":
            state["cart"].append({"item_id": item_id, "quantity": quantity, "unit_price": item.get("price")})
        elif tool_name.startswith("save_"):
            if item_id not in state["saved"]:
                state["saved"].append(item_id)
        elif tool_name == "subscribe_shipment_updates":
            state["subscriptions"].append({"item_id": item_id, "channel": arguments.get("channel", "in_app"), "status": "sandbox_subscribed"})
        elif tool_name == "move_card":
            allowed = {"list-todo": {"list-doing"}, "list-doing": {"list-todo", "list-done"}, "list-done": {"list-doing"}}
            target_list = str(arguments.get("target_list_id") or "")
            if target_list not in allowed.get(str(item.get("list_id") or ""), set()):
                return {"status": "transition_not_allowed", "from_list_id": item.get("list_id"), "target_list_id": target_list}
            item["list_id"] = target_list
            event = {"tool_name": tool_name, "item_id": item_id, "actor_id": arguments.get("actor_id"), "target_list_id": target_list}
            state["activity"].append(event)
            catalog = self.catalog(); catalog["items"] = [item if row.get("id") == item.get("id") else row for row in catalog.get("items", [])]; self._write(self.catalog_path, catalog)
        elif tool_name == "add_comment":
            if not str(arguments.get("actor_id") or "").strip() or not str(arguments.get("comment") or "").strip():
                return {"status": "comment_identity_required", "message": "留言需要已登入使用者與非空白內容。"}
            state["activity"].append({"tool_name": tool_name, "item_id": item_id, "actor_id": arguments.get("actor_id"), "comment": arguments.get("comment")})
        elif tool_name in {"update_checklist", "assign_card", "acknowledge_shipment_exception"}:
            event = {"tool_name": tool_name, "item_id": item_id, "changes": {key: value for key, value in arguments.items() if key not in {"confirmed", "idempotency_key"}}}
            state["activity"].append(event)
        elif action_kind == "transaction":
            item["inventory"] = int(item.get("inventory", 0)) - quantity
            catalog = self.catalog()
            catalog["items"] = [item if row.get("id") == item_id else row for row in catalog.get("items", [])]
            self._write(self.catalog_path, catalog)
            booking_id = "booking_" + secrets.token_hex(6)
            snapshot = {"id": booking_id, "item_id": item_id, "quantity": quantity, "unit_price": item.get("price"), "currency": item.get("currency"), "status": "sandbox_pending_payment"}
            state["bookings"].append(snapshot)
            voucher_id = ""
            if tool_name == "prepare_booking":
                voucher_id = "voucher_" + secrets.token_hex(6)
                state["vouchers"][voucher_id] = {"status": "sandbox_issued", "voucher_id": voucher_id, "booking_id": booking_id, "qr_payload": voucher_id}
            response = {"status": "sandbox_pending_payment", "order": snapshot, "voucher_id": voucher_id or None, "payment_provider_enabled": False}
            state["idempotency"][idempotency_key] = response
            self._write(self.state_path, state)
            return response
        else:
            state["bookings"].append({"item_id": item_id, "status": "requested"})
        response = {"status": "ok", "tool_name": tool_name, "item_id": item_id}
        if idempotency_key:
            state["idempotency"][idempotency_key] = response
        self._write(self.state_path, state)
        return response

    def create_candidate(self, supplier_id: str, rows: list[dict[str, Any]]) -> dict[str, Any]:
        if not supplier_id.strip():
            return {"status": "invalid_supplier", "message": "供應商代號不可空白", "can_publish": False}
        batch_id = "import_" + secrets.token_hex(8)
        errors = []
        normalized = []
        seen = set()
        for index, row in enumerate(rows, 2):
            required_columns = MANIFEST.get("required_columns", ["id", "name", "status", "price", "inventory"])
            missing = [key for key in required_columns if row.get(key) in (None, "")]
            if row.get("id") in seen:
                missing.append("duplicate_id")
            inventory = row.get("inventory")
            if "price" in required_columns or row.get("price") not in (None, ""):
                try:
                    if float(row.get("price", -1)) < 0:
                        missing.append("price")
                except (TypeError, ValueError):
                    missing.append("price")
            if "inventory" in required_columns or inventory not in (None, ""):
                try:
                    inventory = int(inventory)
                    if inventory < 0:
                        missing.append("inventory")
                except (TypeError, ValueError):
                    missing.append("inventory")
            seen.add(row.get("id"))
            if missing:
                errors.append({"row": index, "fields": sorted(set(missing))})
            else:
                normalized_row = dict(row)
                if row.get("price") not in (None, ""):
                    normalized_row["price"] = float(row["price"])
                if inventory not in (None, ""):
                    normalized_row["inventory"] = inventory
                normalized.append(normalized_row)
        candidates = self._read(self.candidates_path, {})
        candidates[batch_id] = {"supplier_id": supplier_id, "rows": normalized, "errors": errors, "status": "candidate"}
        self._write(self.candidates_path, candidates)
        return {"batch_id": batch_id, "status": "candidate", "row_count": len(rows), "valid_count": len(normalized), "error_count": len(errors), "can_publish": not errors}

    def create_candidate_from_file(self, supplier_id: str, filename: str, content: bytes) -> dict[str, Any]:
        suffix = Path(filename).suffix.lower()
        if suffix == ".csv":
            rows = list(csv.DictReader(io.StringIO(content.decode("utf-8-sig"))))
        elif suffix == ".json":
            payload = json.loads(content.decode("utf-8"))
            rows = payload if isinstance(payload, list) else payload.get("items", [])
        elif suffix == ".xlsx":
            from openpyxl import load_workbook
            sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).active
            values = list(sheet.iter_rows(values_only=True))
            headers = [str(value or "").strip() for value in values[0]] if values else []
            rows = [{headers[index]: value for index, value in enumerate(row) if index < len(headers)} for row in values[1:]]
        else:
            return {"status": "unsupported_format", "accepted_formats": [".csv", ".json", ".xlsx"]}
        return self.create_candidate(supplier_id, rows)

    def publish_candidate(self, batch_id: str) -> dict[str, Any]:
        candidates = self._read(self.candidates_path, {})
        batch = candidates.get(batch_id)
        if not batch:
            return {"status": "not_found", "batch_id": batch_id}
        if batch.get("errors"):
            return {"status": "validation_failed", "batch_id": batch_id, "errors": batch["errors"]}
        current = self.catalog()
        by_id = {item["id"]: item for item in current.get("items", [])}
        by_id.update({item["id"]: item for item in batch["rows"]})
        version = "published_" + secrets.token_hex(6)
        self._write(self.catalog_path, {"version": version, "items": list(by_id.values())})
        batch["status"] = "published"
        candidates[batch_id] = batch
        self._write(self.candidates_path, candidates)
        return {"status": "published", "batch_id": batch_id, "version": version, "published_rows": len(batch["rows"])}

    def redeem(self, voucher_id: str, *, confirmed: bool) -> dict[str, Any]:
        if not confirmed:
            return {"status": "confirmation_required", "message": "核銷不可逆，請授權人員再次確認。"}
        state = self._read(self.state_path, {"saved": [], "cart": [], "bookings": [], "vouchers": {}, "idempotency": {}})
        voucher = state["vouchers"].get(voucher_id)
        if not voucher:
            return {"status": "not_found", "voucher_id": voucher_id}
        if voucher.get("status") in {"redeemed", "sandbox_redeemed"}:
            return {"status": "already_redeemed", "voucher_id": voucher_id}
        voucher["status"] = "sandbox_redeemed"
        state["vouchers"][voucher_id] = voucher
        self._write(self.state_path, state)
        return {"status": "sandbox_redeemed", "voucher_id": voucher_id, "live_entitlement": False}

    def _ui(self, state: str, title: str, items: list[dict[str, Any]], pagination: dict[str, Any], *, content_type: str = "list") -> dict[str, Any]:
        normalized = []
        for item in items:
            media = [{"kind": "image", "url": url, "alt": str(item.get("name") or "項目圖片")} for url in item.get("image_urls", []) if str(url).startswith("https://")]
            if str(item.get("video_url", "")).startswith("https://"):
                media.append({"kind": "video", "url": item["video_url"], "alt": str(item.get("name") or "項目影片")})
            normalized.append({"id": item.get("id"), "title": item.get("name"), "summary": item.get("description"), "price_label": f"{item.get('currency', '')} {item.get('price', '')}".strip(), "media": media, "fields": {key: value for key, value in item.items() if (isinstance(value, (str, int, float, bool)) or (MANIFEST.get("runtime_variant") == "travel" and key in {"available_dates", "audience", "languages"} and isinstance(value, list))) and key not in {"description", "name"} and (key != "booking_url" or item.get("booking_link_verified") is True)}, "latitude": item.get("latitude"), "longitude": item.get("longitude")})
        selected_id = str(items[0].get("id")) if items else ""
        actions = [{"id": tool["name"], "label": tool["description"], "kind": tool["action_kind"], "tool_name": tool["name"], "arguments": {"item_id": selected_id} if selected_id else {}, "requires_confirmation": tool["action_kind"] in {"write", "transaction"}} for tool in MANIFEST["tools"]]
        return {"status": "ok", "structuredContent": {"schema_version": "1.0", "state": state, "content_type": content_type, "title": title, "items": normalized, "actions": actions, "pagination": pagination}}
